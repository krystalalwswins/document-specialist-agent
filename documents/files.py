"""CSV/XLSX/text support with explicit size and shape checks."""
import base64
import csv
import hashlib
import io
import json
from pathlib import PurePosixPath
import zipfile

MAX_BYTES = 2 * 1024 * 1024


def safe_name(name):
    if (not name or name in {'.', '..'} or PurePosixPath(name).name != name
            or any(c in name for c in '\\:\x00') or any(ord(c) < 32 for c in name)):
        raise ValueError('filename must be a plain basename')
    if PurePosixPath(name).suffix.lower() not in {'.csv', '.xlsx', '.txt', '.md', '.json'}:
        raise ValueError('supported files: csv, xlsx, txt, md, json')
    return name


def decode_inputs(inputs):
    if len(inputs) > 5:
        raise ValueError('at most 5 input files')
    decoded = []
    names = set()
    total = 0
    for item in inputs:
        name = safe_name(item['filename'])
        if name in names:
            raise ValueError('duplicate input filename')
        names.add(name)
        encoded = item['content_base64']
        if len(encoded) > MAX_BYTES * 2:
            raise ValueError('input too large')
        data = base64.b64decode(encoded, validate=True)
        total += len(data)
        if not data or total > MAX_BYTES:
            raise ValueError('inputs must be nonempty and total at most 2 MiB')
        decoded.append((name, data))
    return decoded


def inspect_document(filename, data, *, required_columns=(), min_rows=1):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('document must be nonempty and at most 2 MiB')
    suffix = PurePosixPath(filename).suffix.lower()
    rows = None
    if suffix == '.csv':
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig')), strict=True))
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError('xlsx expanded size exceeds limit')
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book.active
            if sheet.max_row and sheet.max_row > 10000 or sheet.max_column and sheet.max_column > 200:
                raise ValueError('xlsx dimensions exceed limit')
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if len(rows) >= 10000 or len(row) > 200:
                    raise ValueError('xlsx dimensions exceed limit')
                rows.append(list(row))
        finally:
            book.close()
    if rows is not None:
        rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
        if not rows:
            raise ValueError('table is empty')
        columns = [str(v) if v is not None else '' for v in rows[0]]
        if not all(columns) or len(set(columns)) != len(columns):
            raise ValueError('table headers must be nonempty and unique')
        if any(len(row) != len(columns) for row in rows[1:]):
            raise ValueError('inconsistent table row width')
        if not set(required_columns).issubset(columns) or len(rows) - 1 < min_rows:
            raise ValueError('artifact does not meet required columns/min_rows')
        return {'format': suffix[1:], 'columns': columns, 'row_count': len(rows) - 1,
                'preview': rows[:21]}
    if suffix not in {'.txt', '.md', '.json'}:
        raise ValueError('unsupported document format')
    if required_columns:
        raise ValueError('column requirements need CSV or XLSX')
    text = data.decode('utf-8-sig')
    if not text.strip():
        raise ValueError('document has no content')
    if suffix == '.json':
        json.loads(text)
    return {'format': suffix[1:], 'preview': text[:4000]}


def artifact_metadata(filename, key, data, requirements):
    expected = requirements.get('format')
    if expected and PurePosixPath(filename).suffix.lower() != '.' + expected:
        raise ValueError('artifact format does not match requirement')
    info = inspect_document(filename, data, required_columns=requirements.get('required_columns', []),
                            min_rows=requirements.get('min_rows', 1))
    info.pop('preview', None)
    return {'filename': filename, 'object_key': key, 'size_bytes': len(data),
            'sha256': hashlib.sha256(data).hexdigest(), 'validated': True, **info}
